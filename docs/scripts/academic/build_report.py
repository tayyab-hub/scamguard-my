from pathlib import Path
from copy import deepcopy
import re,json,hashlib,zipfile,shutil
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path.cwd();D=R/'docs';tmp=Path(__file__).resolve().parent/'inputs'
src=Path(json.loads((tmp/'template-contract.json').read_text())['path'])
out=R/'output/documents/SCAMGUARD_CSC1250_FINAL_CANDIDATE_2026-10-20.docx';out.parent.mkdir(parents=True,exist_ok=True)
contract=json.loads((tmp/'template-contract.json').read_text())
assert hashlib.sha256(src.read_bytes()).hexdigest()==contract['sha256']
shutil.copy2(src,out);doc=Document(out);body=doc._element.body
sections=[deepcopy(s._sectPr) for s in doc.sections]
# The cover is an editable slot. Keep the source logo, paragraph styles and page furniture.
for p in doc.paragraphs[:22]:
 p.paragraph_format.line_spacing=p.paragraph_format.line_spacing or 1
 p.paragraph_format.space_after=p.paragraph_format.space_after or Pt(0)
 p.paragraph_format.space_before=p.paragraph_format.space_before or Pt(0)
 if p.text=='CSC1250 Capstone Project':
  p.style='Normal'
  for r in p.runs:r.font.size=Pt(20);r.font.bold=True
 if p.text=='<Proposal Title>':
  p.text='SCAMGUARD: Multi-Modal Scam Detection\nand Reporting Web Application'
  for r in p.runs:r.font.size=Pt(18)
 elif p.text=='PROPOSAL REPORT':
  for r in p.runs:r.text=r.text.replace('PROPOSAL','FINAL')
 elif p.text in ['<Student Name>','<Admission Number>','< DCS or DIT in full>','August 2026']:p.text=''
 elif p.text.startswith('Supervisor :'):p.text='Supervisor : '
cover=[]
for el in list(body):
 if el.tag==qn('w:sectPr'):break
 cover.append(deepcopy(el))
 if el.find('./w:pPr/w:sectPr',el.nsmap) is not None:break
for el in list(body):body.remove(el)
for el in cover:body.append(el)
for sty in doc.styles:
 if sty.type==1:
  sty.font.name='Times New Roman';sty.font.color.rgb=RGBColor(0,0,0)
for name in ['Normal','Body Text']:
 s=doc.styles[name];s.font.size=Pt(12);s.paragraph_format.line_spacing=2;s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;s.paragraph_format.space_after=Pt(0)
for i,size in [(1,20),(2,16),(3,14)]:
 s=doc.styles[f'Heading {i}'];s.font.size=Pt(size);s.font.bold=True;s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.15;s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
for name in ['TOC 1','TOC 2','TOC 3','Table of Figures']:
 if name in doc.styles:
  s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(12);s.paragraph_format.line_spacing=1.3
if 'Caption' not in doc.styles:doc.styles.add_style('Caption',1)
cap=doc.styles['Caption'];cap.font.name='Times New Roman';cap.font.size=Pt(12);cap.font.italic=False;cap.font.color.rgb=RGBColor(0,0,0);cap.paragraph_format.line_spacing=1.25
def clean(s):
 s=re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',s);s=s.replace('**','').replace('`','').replace('*','')
 return s.replace(' → ',' → ').strip()
def p(text='',style=None):
 x=doc.add_paragraph(clean(text),style or 'Normal');x.paragraph_format.widow_control=True;return x
def h(text,level=2):
 x=p(text,f'Heading {level}');num=OxmlElement('w:numPr');ni=OxmlElement('w:numId');ni.set(qn('w:val'),'0');num.append(ni);x._p.get_or_add_pPr().append(num)
def endsec(i):
 x=doc.add_paragraph();x.add_run().font.size=Pt(1);x.paragraph_format.space_after=Pt(0);x.paragraph_format.line_spacing=Pt(1);x.paragraph_format.keep_with_next=False;x._p.get_or_add_pPr().append(deepcopy(sections[i]))
def field(par,code,cache=''):
 run=par.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'begin');run._r.append(a)
 r=par.add_run();i=OxmlElement('w:instrText');i.set(qn('xml:space'),'preserve');i.text=' '+code+' ';r._r.append(i)
 r=par.add_run();sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep);par.add_run(cache)
 r=par.add_run();e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end');r._r.append(e)
captions=[]
def caption(kind,n,title,source,above=False):
 x=doc.add_paragraph(style='Caption');x.add_run(kind+' ');field(x,'SEQ '+kind+' \\* ARABIC',str(n));x.add_run('. '+title);x.paragraph_format.keep_with_next=True
 captions.append((kind,n,title,source));s=p('Source: '+source);s.paragraph_format.line_spacing=1.1;s.paragraph_format.space_after=Pt(6);s.paragraph_format.keep_with_next=above
def table(n,title,headers,rows,widths=None,source='Project source and evidence at final release; see Appendix B.'):
 phrase=title if title.split()[0].isupper() else title[0].lower()+title[1:]
 p(f'Table {n} summarises the '+phrase.replace('gantt','Gantt')+'.')
 caption('Table',n,title,source,True);t=doc.add_table(rows=1, cols=len(headers));t.autofit=False
 # Retain formal monochrome table styling; content wraps, header repeats, rows stay intact.
 for j,txt in enumerate(headers):t.rows[0].cells[j].text=clean(txt)
 for row in rows:
  for j,txt in enumerate(row):t.add_row() if False else None
  cells=t.add_row().cells
  for j,txt in enumerate(row):cells[j].text=clean(str(txt))
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');pr.append(ns)
  if ri==0:
   repeat=OxmlElement('w:tblHeader');pr.append(repeat)
  for j,c in enumerate(row.cells):
   if widths:c.width=Inches(widths[j])
   tcpr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'999999');b.append(el)
   tcpr.append(b)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDEB');tcpr.append(sh)
   for par in c.paragraphs:
    par.paragraph_format.line_spacing=2;par.paragraph_format.space_after=Pt(0);par.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;par.paragraph_format.keep_with_next=(ri==0)
    for r in par.runs:r.font.name='Times New Roman';r.font.size=Pt(12);r.bold=(ri==0)
 gap=p();gap.add_run().font.size=Pt(1);gap.paragraph_format.line_spacing=Pt(1);gap.paragraph_format.space_after=Pt(6)
 return t
def fig(n,title,file,source,width=6.0):
 phrase=title if title.split()[0].isupper() else title[0].lower()+title[1:]
 intro=p(f'Figure {n} illustrates '+phrase+'.');intro.paragraph_format.keep_with_next=True
 x=doc.add_paragraph();x.paragraph_format.line_spacing=1;x.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.keep_with_next=True
 shape=x.add_run().add_picture(str(D/'figures'/file),width=Inches(width));shape._inline.docPr.set('descr',title)
 caption('Figure',n,title,source)
def blocks(name,skip_intro=False):
 text=(D/name).read_text(encoding='utf-8');chunks=re.split(r'\n\s*\n',text)
 for c in chunks:
  if c.startswith('#') or c.startswith('|') or c.startswith('**Table') or c.startswith('```'):continue
  if c.startswith('References:') or (skip_intro and c==chunks[1]):continue
  p(c.replace('\n',' '))
def mdtable(name,index=0):
 s=(D/name).read_text(encoding='utf-8');groups=re.findall(r'(?m)^\|.*(?:\n\|.*)+',s);lines=groups[index].splitlines();return [[v.strip() for v in line.strip('|').split('|')] for line in lines[2:]]
def front(title,code,i):
 x=p(title);x.runs[0].bold=True;x.runs[0].font.size=Pt(20);field(p(),code);endsec(i)
front('Table of Contents','TOC \\o "1-2" \\h \\z \\u',1)
front('List of Figures','TOC \\h \\z \\c "Figure"',2)
front('List of Tables','TOC \\h \\z \\c "Table"',3)
h('1 Introduction',1)
p('SCAMGUARD is an authenticated web application for inspecting unfamiliar messages, URLs, telephone numbers and QR payloads. The formal project title retains “Reporting”; the implemented reporting scope is private assessment history. The application provides decision support with explicit evidence limits. This chapter establishes the user problem and translates it into the four proposal objectives and bounded deliverables.')
h('1.1 Problem statement')
for para in re.split(r'\n\s*\n',(D/'FINAL_PROBLEM_STATEMENT.md').read_text(encoding='utf-8'))[1:5]:p(para)
h('1.2 Objectives')
p('The user-designated proposal supplies four objectives: identify requirements, design the system, develop and deploy, and evaluate through functional tests, classifier evaluation and a volunteer usability study. Table 1 preserves that wording and the proposed phases while distinguishing present evidence from planned work. O4 is partially complete: the study has not yet been conducted. The proposal is evidence of intended scope and dates; signed approval and the final assessment timetable are separate.')
objectives=mdtable('FINAL_SMART_OBJECTIVES.md')
table(1,'SMART objectives and acceptance criteria',['ID','Specific outcome, measurement and status'],[[r[0],r[1]+'\nOutcome: '+r[2]+'\nMeasure: '+r[3]+'\nEvidence: '+r[4]+'\nPhase: '+r[5]+'\nStatus: '+r[6]] for r in objectives],[.65,5.62])
trace=mdtable('FINAL_SMART_OBJECTIVES.md',1)
table(2,'Objective-to-evidence traceability',['ID','Implementation, verification and evaluation'],[[r[0],r[1]+'\nVerification: '+r[2]+'\nEvaluation: '+r[3]+'\n'+r[4]] for r in trace],[.65,5.62])
h('1.3 Scope')
for para in re.split(r'\n\s*\n',(D/'FINAL_SCOPE_DELIVERABLES.md').read_text(encoding='utf-8'))[1:4]:p(para)
p('Out of scope are community reporting and moderation, live URL fetching, caller identification, reputation feeds, general OCR, payment authorisation, complete payment-scheme validation and a combined cross-modal fraud probability. Optional external AI review is backend-only, disabled by default and excluded from the recorded evaluation.')
h('1.4 Deliverables')
scope=mdtable('FINAL_SCOPE_DELIVERABLES.md')
table(3,'Deliverables and problem contributions',['Deliverable','Description, contribution and verification'],[[r[0],r[1]+'\nContribution: '+r[2]+'\nEvidence: '+r[3]+'\nStatus: '+r[4]] for r in scope],[1.35,4.92])
h('1.5 Chapter organisation')
p('Chapter 2 compares three related systems. Chapter 3 explains the solution and justifies its requirements and resources. Chapter 4 separates the proposal plan from observed Git milestones and risks. Appendices A–D connect design, evaluation and the illustrated user workflow to limitations and the conclusion. Unperformed human validation remains distinct from technical results.')
endsec(4)
h('2 Background of study',1)
for c in re.split(r'\n\s*\n',(D/'BACKGROUND_OF_STUDY.md').read_text(encoding='utf-8'))[1:]:
 if c.startswith('## '):h({'VirusTotal':'2.1 VirusTotal','Google Safe Browsing':'2.2 Google Safe Browsing','Truecaller':'2.3 Truecaller','Synthesis and design implications':'2.4 Synthesis and design implications'}[c[3:]])
 else:p(c)
comp=mdtable('RELATED_SYSTEM_COMPARISON.md')
table(4,'Related-system scope, design and implementation',['System / users','Critical comparison'],[[r[0],'Scope/deliverables: '+r[1]+'\nDesign/implementation: '+r[2]+'\nStrength: '+r[3]+'\nTrade-off: '+r[4]] for r in comp],[1.25,5.02],source='VirusTotal (n.d.-a, n.d.-b, n.d.-c), Google (2026a, 2026b), Truecaller (n.d.-a, n.d.-b, n.d.-c); SCAMGUARD release source.')
caps=mdtable('RELATED_SYSTEM_COMPARISON.md',1)
table(5,'Documented capability evidence',['System','Capability and explanation boundary'],[[r[0],'Message: '+r[1]+'; URL: '+r[2]+'; Phone: '+r[3]+'; QR: '+r[4]+'.\nExplanation: '+r[5]+'\nPrivate history: '+r[6]+'\nCamera: '+r[7]] for r in caps],[1.25,5.02],source='Same primary sources as Table 4. Y = documented; NE = NOT YET EVIDENCED in reviewed scope, not absent.')
endsec(5)
h('3 Resource approach',1);h('3.1 Solution approach')
for c in re.split(r'\n\s*\n',(D/'SOLUTION_APPROACH.md').read_text(encoding='utf-8'))[1:]:
 if c.startswith('## '):h(c[3:],3)
 elif not c.startswith('```'):p(c)
fig(1,'Production architecture and service responsibilities','figure-01-architecture.png','Final release architecture; provider routing evidence in PRODUCTION_ACCEPTANCE.md.',5.6)
fig(2,'Analysis data flow and retention boundary','figure-02-data-flow.png','Analysis services and privacy model at final release.',5.5)
fig(3,'Selected database entities and ownership relationships','figure-03-erd.png','backend/app/db/models.py; Alembic head 0006_qr_intelligence.',5.5)
fig(4,'Authentication, protected access and recovery','figure-04-authentication.png','backend/app/core/auth.py, auth API and password-reset services.',5.5)
fig(5,'QR upload, camera consent and payload routing','figure-05-qr-workflow.png','QR engine/decoder/payment modules and frontend cameraScanner.ts.',5.5)
h('3.2 Software and hardware requirements')
p('Tables 6–7 assign acceptance IDs to the implemented requirements. Each requirement maps to an objective, implementation path and verification path in REQUIREMENTS.md. The compact report tables retain the behaviour and evidence boundary; the companion matrix contains the full paths. A status of locally verified does not establish unmeasured production or human outcomes.')
data=json.loads((tmp/'tables.json').read_text())
for row in data['fr']:row[2]='O2/O3; evaluated under O4'
for row in data['nfr']:
 if row[0]=='NFR-06':row[2]=row[2].replace('320\uFFFD1440','320–1440')
 if row[0]=='NFR-07':row[4]='docs/evidence/completion-verification.json'
 if row[0]=='NFR-09':row[4]='docs/evidence/completion-message-evaluation.json'
 if row[0]=='NFR-10':
  row[2]='Serve the verified runtime source on Vercel/Render with ready database and current Alembic head; path-scoped source equivalence may explain later documentation-only commits.'
  row[4]='docs/FINAL_COMPLETION_AUDIT_2026-10-20.md'
  row[5]='Observed 20–21 October 2026; point-in-time evidence, not continuous monitoring.'
table(6,'Functional requirements',['ID / objective','Requirement and verification boundary'],[[r[0]+'\n'+r[2],r[1]+'\nEvidence: '+Path(r[4]).name+'\n'+r[5]] for r in data['fr']],[1,5.27])
table(7,'Non-functional requirements',['ID / quality','Acceptance and evidence boundary'],[[r[0]+'\n'+r[1],r[2]+'\nImplementation: '+r[3]+'\nVerification: '+r[4]+'\n'+r[5]] for r in data['nfr']],[1.25,5.02])
p('All technologies in Table 8 are represented by repository dependencies or deployed services. Versions refer to the release, not upgrade recommendations. Alternatives are retrospective engineering comparisons except the Message/URL candidates actually evaluated. Hardware roles in Table 9 distinguish developer, client and managed server resources; no purchased-server specification or unsupported minimum device is invented.')
table(8,'Software purposes, selection reasons and trade-offs',['Technology','Justification and alternative'],[[r[0],'Purpose: '+r[1]+'\nRequired because: '+r[2]+'\nSelection: '+r[3]+'\nAlternative: '+r[4]+'\nTrade-off: '+r[5]] for r in data['software']],[1.4,4.87],source='Release dependency declarations and locks; Google (n.d.), MDN contributors (n.d.), OWASP Foundation (n.d.-a, n.d.-b), Render (n.d.), Vercel (n.d.).')
table(9,'Hardware and environment roles',['Role','Requirement and justification'],[[r[0],'\n'.join(r[1:])] for r in data['hardware']],[1.4,4.87])
h('3.3 Detailed development planning')
p('Development proceeded through bounded tasks: establish the application and persistence first, then add reproducible local intelligence, account ownership and QR routing. Later tasks integrated camera lifecycle and retrieval controls. Each phase produced a reviewable deliverable with regressions. The final freeze permitted integration corrections and evaluation, including two QR credential-redaction edge cases, while preserving model methodology and the migration head.')
p('Model work separated preparation, validation selection and held-out reporting. Security work used isolated local PostgreSQL databases and synthetic accounts. Deployment closure gated the normal main merge on a full verification rerun, followed by matching provider identities and read-only production probes. Chapter 4 records the proposed 1 September–20 November calendar and observed milestones separately. Proposed report/presentation dates are 18/20 November, subject to the official timetable.')
endsec(6)
h('4 Project planning and risk assessment',1);h('4.1 Observed development chronology')
p('The chronology below uses Git timestamps in Asia/Kuala_Lumpur. A first commit records an artifact, not necessarily the day work began. The repository starts with a completed foundation on 3 September 2026. Dates therefore establish observed commit coverage only; they do not establish labour duration, original approval dates or whether a milestone met an earlier deadline.')
plan=json.loads((tmp/'plan.json').read_text());phases=plan['phases']
table(10,'Recorded phases, activities and milestones',['Phase / dates','Tasks, deliverable and milestone'],[[r[0]+'\n'+r[5]+' to '+r[6],r[1]+'\nActivities: '+r[2]+'\nDeliverable: '+r[3]+'\nMilestone: '+r[4]+'\nStatus: '+r[7]] for r in phases],[1.65,4.62],source='evidence/project-git-chronology.txt; complete commit identities are retained in that record.')
h('4.2 Gantt chart and deadlines')
fig(6,'Proposed objective phases and current evidence','completion-plan.png','User-designated proposal, PDF pages 7 and 43–47; planned dates are not actual completion.',6.15)
table(11,'Proposed milestones and evidence boundary',['Milestone','Planned date / present qualification'],[['M1 proposal acceptance','11 September; approval event not independently evidenced'],['M2 design review','18 September; design artifacts available, review date not inferred'],['M3 deployment','16 October; separate September/October deployment records exist'],['M4 feature freeze','2 November; planned checkpoint, no new feature phase'],['M5 report','By 18 November; official deadline to confirm'],['M6 presentation','By 20 November; official date and duration to confirm']],[2.2,4.07],source='User-designated proposal, Table 14, PDF page 45. All dates are 2026.')
p('The proposal supplies a 12-week plan from 1 September to 20 November 2026. O1 spans Weeks 1–3, O2 Weeks 2–3, O3 Weeks 3–9 and O4 Weeks 9–11. Table 10 is historical Git evidence, while Figure 6 and Table 11 are planned dates. Git intervals are not labour duration or proof of supervisor approval. Actual study, literature-review and rehearsal dates must come from real records; no study completion is invented.')
h('4.3 Risk assessment')
p('Risk likelihood and impact use a project judgement scale from 1 (low) to 3 (high). Their product gives low 1–2, medium 3–4 and high 6–9. These are prioritisation judgements, not measured incident probabilities. Table 12 links each risk to a specific implemented mitigation, contingency and residual limitation.')
table(12,'Risk register and contingencies',['Risk / category / level','Mitigation, contingency and residual status'],[[r[0]+' '+r[1]+'\n'+r[2]+'\nL='+r[3]+'; I='+r[4]+'\n'+r[5],'Mitigation: '+r[6]+'\nContingency: '+r[7]+'\nResidual: '+r[8]+'\nStatus: '+r[9]] for r in plan['risks']],[1.65,4.62],source='Project risk judgement tied to recorded controls; full ten-column register in FINAL_RISK_REGISTER.md.')
# The preceding risk register states residual risk; the evaluation connection is already explicit.
endsec(7)
h('5 References',1)
refs=re.split(r'\n\s*\n',(D/'FINAL_REFERENCES.md').read_text(encoding='utf-8'))[2:-1]
for ref in refs:
 x=p()
 for i,part in enumerate(ref.split('*')):x.add_run(part).italic=(i%2==1)
 x.paragraph_format.left_indent=Inches(.35);x.paragraph_format.first_line_indent=Inches(-.35);x.paragraph_format.space_after=Pt(6)
endsec(8)
h('Appendix A — Design and trust boundaries',1)
p('Figures 3–5 in Chapter 3 expand the design and trust boundaries. They are explanatory views derived from the released code, not a claim of an independently certified security architecture. The entity-relationship view shows selected fields; the actual schema contains additional result and version metadata.')
p('EMVCo (n.d.) distinguishes merchant-presented and consumer-presented payment QR uses from general QR content. SCAMGUARD implements a generic structural and CRC subset, not full compliance with either payment scheme or proof of recipient legitimacy.')
p('Production origin/CSRF guards and session-derived user identity protect mutation paths; owner-scoped database queries protect records independently of interface routing. Original uploads and camera frames are not retained. Specified QR credentials are removed before new stored records are created, while original-byte assessment remains unchanged. No retroactive rewrite of old production rows is claimed.')
endsec(9)
h('Appendix B — Verification and evaluation',1);h('B.1 Release identity and evidence')
release=json.loads((D/'evidence/completion-release.json').read_text())
p('This candidate incorporates the October v2 implementation and the 20 October completion checks. Application merge: '+release['application_merge_sha']+'. Vercel and Render source identities and dated production observations are recorded in FINAL_COMPLETION_AUDIT_2026-10-20.md. The expected migration head remains 0006_qr_intelligence, confirmed by a SELECT-only production observation. This pass changes dependency pins and adds authored regressions; trained artifacts, data splits, scoring policies and migrations remain unchanged.')
p('The complete local gate recorded 365 Pytest passes, 164 Vitest passes and 52 Playwright passes. One paid/live external-AI test is intentionally skipped. Backend compatibility and persisted browser checks were repeated after the urllib3 security patch. Machine records distinguish these runs and preserve warnings and initial environment failures. They do not establish continuous availability or independent penetration testing.')

tests=mdtable('FINAL_EVALUATION.md')
table(13,'October completion verification',['Check','Recorded result and interpretation'],[[r[0],r[1]+'\n'+r[2]] for r in tests],[1.45,4.82])
h('B.2 Learned components and controlled cases')
p('Message data originate from Mishra and Soni (2022), and URL data from Prasad and Chandra (2024), both published under CC BY 4.0. The Message source has 5,971 rows; cleaning retains 5,797 after duplicate/conflicting rows are removed. Near-duplicate grouping precedes the 3,477/1,160/1,160 train/validation/test split. Word unigram/bigram TF–IDF and Logistic Regression were selected on validation macro F1 against SVM and Naive Bayes. The selected model was refitted on train plus validation before the frozen test was evaluated.')
p('The URL source has 235,795 rows and the processed corpus 234,674. Registrable-domain grouping produces 164,074 training, 33,699 validation and 36,901 held-out rows. The exported 120-tree forest uses 27 features derived locally from URL text. All legitimate training URLs are HTTPS homepages without queries; this strong collection bias restricts the apparent held-out performance. Domain separation does not establish temporal or campaign independence.')
evals=mdtable('FINAL_EVALUATION.md',1)
p('The Message and URL matrices were reproduced during the October completion checks. The controlled URL, Phone, QR/payment cases and warm-latency values in Table 14 are historical measurements from 12 September 2026; they were not remeasured for this candidate.')
table(14,'Intelligence evaluation and interpretation',['Evaluation / sample','Result and boundary'],[[r[0]+'\n'+r[1],r[2]+'\n'+r[3]] for r in evals],[2,4.27])
fig(7,'Frozen Message and URL confusion matrices','figure-07-confusion-matrices.png','MODEL_EVALUATION.md, URL_MODEL_EVALUATION.md and frozen reproduction records.',6.15)
fig(8,'October automated test counts','completion-tests.png','completion-verification.json; 365 backend, 164 frontend and 52 browser tests; one live-AI skip.',6.1)
p('Accuracy measures the proportion of all test predictions that match their labels. Precision is the fraction of predictions for a class that are correct; recall is the fraction of actual class members recovered. F1 is their harmonic mean, and macro F1 weights each class equally (scikit-learn developers, n.d.). For Message, 13 of 106 SCAM test examples were predicted SPAM. Its SCAM recall is 93/106 = 87.7358%, despite overall accuracy of 96.8966%. URL phishing recall is 16,495/16,966 = 97.2239%. These are classifier metrics, not measurements of the final rules/model fusion or real-world fraud prevention.')
p('The four proposal objectives have concrete evidence mappings in Table 2. O4 remains partial until the proposed volunteer study is performed. All 18 functional and 10 non-functional requirements have implementation and verification references in REQUIREMENTS.md. Requirements involving physical cameras, human usability/accessibility or inbox delivery retain qualified statuses. A recorded test is evidence for its asserted behaviour, not automatic completion of every real-world outcome.')
h('B.3 Human evaluation protocol')
p('No formal volunteer study has been conducted. The five core tasks follow the proposal: account access, Message analysis, URL analysis, risk/explanation interpretation and History retrieval. QR, Phone and payment interpretation are optional supplementary probes. With appropriate institutional approval and participant consent, record pseudonym, device, date, unassisted completion, errors, time, explanation of risk versus confidence and optional ease comments. Do not include credentials or sensitive submissions. Report actual sample size, recruitment limits and failures. Blank forms and the protocol are in USABILITY_TEST_PLAN.md and USABILITY_RESULTS_TEMPLATE.md; no participant count or score is claimed here.')
endsec(10)
h('Appendix C — User workflow and screenshots',1)
p('Figures 9–11 are fresh authenticated production captures from the supplementary completion capture run using only synthetic accounts and authored fixtures. They illustrate the deployed interface; they are not a usability study. Capture provenance, source identity and cleanup are recorded in completion-production-captures.json, completion-production-smoke.json and the final audit. Captured views remain available as the offline demonstration fallback. No existing user data were accessed.')
fig(9,'Message assessment with separate risk and confidence','completion-message.png','Completion production smoke; authored password/OTP-request fixture.',6.15)
p('Enter permitted message text, select Analyse and inspect the indicators and limitations. The displayed fusion score is not a probability of fraud. URL mode analyses the supplied string without visiting its destination. Phone mode reports numbering metadata and leaves ordinary valid numbers at Insufficient Evidence.')
fig(10,'QR analysis after an explicit upload submission','completion-qr.png','Completion production smoke; generated reserved-domain QR fixture.',6.15)
p('Choose one supported image and select Analyse QR. For camera input, explicitly start the camera, allow local detection and review the stopped preview before selecting Analyse. Cancellation and navigation stop camera resources. A decodable payload or valid payment CRC does not establish a legitimate destination or recipient; use upload if camera access is unavailable.')
fig(11,'Private History with search, type and risk filters','completion-history.png','Completion production smoke; synthetic records created in the same session.',6.15)
p('History supports owner-scoped retrieval and deterministic ordering. Open a result to revisit its original evidence. Analyse again copies owned Message/URL/Phone content into a new draft; the earlier result remains immutable. QR requires fresh input. Account controls edit profile fields, and recovery uses a single-use expiring link. Production A/B isolation was exercised with two synthetic accounts; the owner checklist still requires controlled-inbox delivery and physical-camera evidence.')
endsec(11)
h('Appendix D — Limitations, further work and conclusion',1);h('D.1 Limitations')
for c in re.split(r'\n\s*\n',(D/'LIMITATIONS.md').read_text(encoding='utf-8')):
 if c.startswith('- '):
  for item in re.split(r'\n(?=- )',c):p(item.removeprefix('- ').replace('\n',' '))
p('The production recovery checklist checks sender restrictions before the demonstration. The default resend.dev test sender is restricted to the Resend account owner’s address; broader recipients require a verified sending domain (Resend, n.d.). This provider rule does not establish the actual configured sender or inbox delivery for this project.')
h('D.2 Prioritised further work')
p('Further validation should precede expansion. Priorities are independent user comprehension and accessibility work, fresh licensed temporal/multilingual model evaluation, calibrated confidence studies, and scoped security review. Operations need measured recovery, inbox delivery, load and backup-restore evidence. Any reputation integration requires provenance, consent, provider-failure behaviour and cost controls. Payment support would need scheme-specific mandatory fields and authoritative test vectors. None of these prospective items is reported as implemented by this academic pass.')
h('D.3 Conclusion')
blocks('FINAL_CONCLUSION.md')
h('D.4 Evidence completion record')
p('MANUAL ACTION REQUIRED: complete the cover identity fields; confirm the final rubric, official timetable and presentation duration; conduct and report the proposed volunteer study; record physical-device and controlled-inbox acceptance; and perform the rehearsals. The proposal resolves the intended scope and planned calendar. Technical checks cannot substitute for human evaluation or oral performance. FINAL_PERSONAL_ACTIONS.md lists the exact remaining owner work. Verdict: TECHNICALLY READY, OWNER ACCEPTANCE PENDING; this is a submission candidate, not a claim of completed assessment requirements.')
p('Abbreviations: API, application programming interface; CRC, cyclic redundancy check; CSRF, cross-site request forgery; ERD, entity-relationship diagram; FR/NFR, functional/non-functional requirement; ML, machine learning; QR, quick response; SHA, commit identifier; TF–IDF, term frequency–inverse document frequency; TLV, tag-length-value; URL, uniform resource locator.')
body.append(deepcopy(sections[12]));doc.core_properties.title='SCAMGUARD — CSC1250 Final Candidate, October 2026';doc.core_properties.subject='Updated final candidate using the supplied proposal template; final applicability requires confirmation';doc.core_properties.author='';doc.core_properties.last_modified_by=''
settings=doc.settings.element;upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');settings.append(upd)
doc.save(out)
# Preserve unaffected source package parts byte-for-byte before native field refresh.
mutable={'word/document.xml','word/styles.xml','word/settings.xml','word/_rels/document.xml.rels','[Content_Types].xml','docProps/core.xml'}
with zipfile.ZipFile(src) as z:original={n:z.read(n) for n in z.namelist()}
with zipfile.ZipFile(out) as z:built={n:z.read(n) for n in z.namelist()}
for name,blob in original.items():
 if name not in mutable:built[name]=blob
with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
 for name,blob in built.items():z.writestr(name,blob)
(tmp/'report-captions.json').write_text(json.dumps(captions,indent=2),encoding='utf-8')
print(f'Created {out}; {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, {len(doc.sections)} sections.')
